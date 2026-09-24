#!/usr/bin/env python3
"""Build Manan's THOA report as PDF + DOCX from manan_report.md (single source of truth)."""
import re
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
MD_PATH = ROOT / "manan_report.md"
OUT_DIR = ROOT / "reports"
FIG_DIR = OUT_DIR / "figures"
LOGO_PATH = OUT_DIR / "vit_logo.png"

BODY_FONT = "Georgia"
BODY_BOLD = "Georgia-Bold"
BODY_ITAL = "Georgia-Italic"
FONT_SIZE = 12
LEADING = round(12 * 1.18, 2)

class Block:
    def __init__(self, kind, text):
        self.kind = kind
        self.text = text

def parse_md() -> tuple:
    raw = MD_PATH.read_text(encoding="utf-8")
    blocks = []
    for chunk in raw.split("[PAGE BREAK]"):
        for line in chunk.splitlines():
            line = line.rstrip()
            if not line.strip() or line.strip() == "---":
                continue
            if line.startswith("[FIGURE:"):
                blocks.append(Block("fig", line[9:-1].strip()))
            elif line.startswith("### "):
                blocks.append(Block("h3", line[4:]))
            elif line.startswith("## "):
                blocks.append(Block("h2", line[3:]))
            elif line.startswith("- "):
                blocks.append(Block("bullet", line[2:]))
            else:
                blocks.append(Block("body", line))
        blocks.append(Block("brk", ""))
    if blocks and blocks[-1].kind == "brk":
        blocks.pop()

    body_start = next(i for i, b in enumerate(blocks)
                      if b.kind == "h2" and b.text.startswith("Abstract"))
    cover = [b for b in blocks[:body_start] if b.kind == "body"]
    body = blocks[body_start:]
    return cover, body

def fill_cover_text(t: str) -> str:
    return t.replace("[VIT LOGO — top center]", "@LOGO@")

def html_escape(t: str) -> str:
    t = t.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    t = re.sub(r"\*\*([^*]+)\*\*", r"<b>\1</b>", t)
    t = re.sub(r"(?<!\*)\*([^*]+)\*(?!\*)", r"<i>\1</i>", t)
    return t

# ----------------------------------------------------------------------------
# Flowcharts: drawn with Pillow -> PNG
# ----------------------------------------------------------------------------
from PIL import Image, ImageFont, ImageDraw

SCALE = 2

_FONT_CACHE = {}

def _font(size, bold=False):
    key = (size, bold)
    if key in _FONT_CACHE:
        return _FONT_CACHE[key]
    candidates = [
        ("/Library/Fonts/Times New Roman.ttf" if not bold else "/Library/Fonts/Times New Roman Bold.ttf", None),
        ("/System/Library/Fonts/Helvetica.ttc", 0 if not bold else 1),
        ("/System/Library/Fonts/Supplemental/Arial.ttf", None),
        ("/System/Library/Fonts/Supplemental/Arial Bold.ttf", None),
    ]
    for path, idx in candidates:
        try:
            f = ImageFont.truetype(path, size * SCALE, index=idx)
            _FONT_CACHE[key] = f
            return f
        except Exception:
            continue
    f = ImageFont.load_default(size * SCALE)
    _FONT_CACHE[key] = f
    return f

BLUE = "#fce7f3"; BLUE_S = "#ec4899"
GREY = "#ede9fe"; GREY_S = "#7c3aed"
GREEN = "#d1fae5"; GREEN_S = "#059669"
AMBER = "#fed7aa"; AMBER_S = "#ea580c"
PALE = "#fefce8"; PALE_S = "#a16207"
INK = "#111827"; LINE = "#374151"

class Canvas:
    def __init__(self, w, h):
        self.w, self.h = w, h
        self.img = Image.new("RGB", (w * SCALE, h * SCALE), "white")
        self.d = ImageDraw.Draw(self.img)

    def box(self, cx, cy, w, h, text, fill=GREY, stroke=GREY_S, size=8, bold=True):
        x0, y0 = (cx - w / 2) * SCALE, (cy - h / 2) * SCALE
        x1, y1 = (cx + w / 2) * SCALE, (cy + h / 2) * SCALE
        self.d.rounded_rectangle([x0, y0, x1, y1], radius=0,
                                 fill=fill, outline=stroke, width=max(2, SCALE))
        lines = text.split("\n")
        lh = size * SCALE * 1.25
        start = (cy * SCALE) - (lh * (len(lines) - 1)) / 2
        for i, ln in enumerate(lines):
            self.d.text((cx * SCALE, start + i * lh), ln, font=_font(size, bold),
                        fill=INK, anchor="mm")

    def arrow(self, x1, y1, x2, y2):
        import math
        s, l = SCALE, 7 * SCALE
        self.d.line([x1 * s, y1 * s, x2 * s, y2 * s], fill=LINE, width=max(2, s))
        ang = math.atan2(y2 - y1, x2 - x1)
        self.d.polygon([(x2 * s, y2 * s),
                        ((x2 - l * math.cos(ang - 0.4)) * s, (y2 - l * math.sin(ang - 0.4)) * s),
                        ((x2 - l * math.cos(ang + 0.4)) * s, (y2 - l * math.sin(ang + 0.4)) * s)],
                       fill=LINE)

    def label(self, x, y, text, size=8, color=INK, bold=True, anchor="mm"):
        self.d.text((x * SCALE, y * SCALE), text, font=_font(size, bold), fill=color, anchor=anchor)

    def note(self, y, text, h=22, fill=PALE, stroke=PALE_S):
        self.box(220, y, 440, h, text, fill=fill, stroke=stroke, size=7.5)

    def save(self, path):
        self.img.save(path)

def fig_pipeline_tech():
    """Figure 1: Proposed end-to-end pipeline."""
    c = Canvas(440, 300)
    c.box(73, 262, 130, 34, "Document upload", fill=PALE, stroke=PALE_S)
    c.arrow(138, 262, 168, 262)
    c.box(220, 262, 150, 34, "Multimodal LLM\nextraction agent", fill=BLUE, stroke=BLUE_S)
    c.arrow(293, 262, 322, 262)
    c.box(367, 262, 146, 34, "Pydantic schema\nvalidation", fill=BLUE, stroke=BLUE_S)
    c.arrow(367, 245, 367, 232)
    c.box(367, 214, 146, 36, "Deterministic rule engine:\n16 rules cited to THOA sections", fill=GREY, stroke=GREY_S)
    c.arrow(367, 196, 367, 184)
    c.box(367, 166, 146, 36, "Flags + explanations\n(no score, no verdict)", fill=GREY, stroke=GREY_S)
    c.arrow(293, 166, 240, 166)
    c.box(165, 166, 150, 36, "Human Authorization\nCommittee decision", fill=GREEN, stroke=GREEN_S)
    c.arrow(90, 148, 90, 130)
    c.box(90, 112, 150, 36, "Immutable audit log", fill=AMBER, stroke=AMBER_S)
    c.label(293, 118, "No LLM anywhere in the decision path; the\nsystem is structurally incapable of approving",
            size=7.5, color="#4a3728", anchor="lm")
    c.note(76, "AI does extraction only. The decision path is deterministic, traceable, and human-terminated.", h=24)
    return c

def fig_ai_boundary():
    """Figure 2: AI/deterministic/human responsibility split."""
    c = Canvas(440, 280)
    c.box(220, 258, 420, 44, "AI EXTRACTION LAYER\nMultimodal LLM (GPT-4o/Claude) + spaCy/BERT NER\nReads documents, extracts structured fields, never decides",
          fill=BLUE, stroke=BLUE_S, size=8)
    c.arrow(220, 236, 220, 222)
    c.box(220, 204, 420, 36, "Pydantic schema validation\nCatches malformed data before decision logic",
          fill=PALE, stroke=PALE_S, size=8)
    c.arrow(220, 186, 220, 172)
    c.box(220, 154, 420, 36, "DETERMINISTIC DECISION LAYER\n16 rules mapped to THOA sections, external config, zero LLM calls",
          fill=GREY, stroke=GREY_S, size=8)
    c.arrow(220, 136, 220, 122)
    c.box(220, 104, 420, 36, "HUMAN REVIEW LAYER\nAuthorization Committee decides; system never approves",
          fill=GREEN, stroke=GREEN_S, size=8)
    c.arrow(220, 86, 220, 72)
    c.box(220, 54, 420, 34, "Immutable audit log: every action traced, every decision logged",
          fill=AMBER, stroke=AMBER_S, size=8)
    c.label(30, 278, "AI", size=9, color=BLUE_S, anchor="lm")
    c.label(30, 204, "", size=9, color=GREY_S, anchor="lm")
    c.label(30, 154, "Rules", size=9, color=GREY_S, anchor="lm")
    c.label(30, 104, "Human", size=9, color=GREEN_S, anchor="lm")
    return c

def fig_ocr_detail():
    """Figure 3: Proposed OCR extraction pipeline detail."""
    c = Canvas(440, 330)
    c.box(220, 310, 180, 30, "Scanned PDF / Image upload", fill=PALE, stroke=PALE_S, size=8)
    c.arrow(220, 295, 220, 282)
    c.box(220, 268, 180, 28, "PDF to image (pdf2image)", fill=BLUE, stroke=BLUE_S, size=8)
    c.arrow(220, 254, 220, 242)
    c.box(220, 228, 200, 28, "OpenCV preprocess\n(deskew, denoise, threshold)", fill=BLUE, stroke=BLUE_S, size=8)
    c.arrow(220, 214, 220, 202)
    c.box(220, 188, 180, 28, "Tesseract OCR\nraw text extraction", fill=BLUE, stroke=BLUE_S, size=8)
    c.arrow(220, 174, 220, 162)
    c.box(220, 148, 200, 28, "Regex + spaCy NER\nfield parsing", fill=BLUE, stroke=BLUE_S, size=8)
    c.arrow(220, 134, 220, 122)
    c.box(220, 108, 200, 28, "Pydantic schema\nvalidation", fill=GREY, stroke=GREY_S, size=8)
    c.arrow(220, 94, 220, 82)
    c.box(220, 68, 200, 28, "Structured JSON\nfed to rule engine", fill=GREY, stroke=GREY_S, size=8)
    c.label(360, 310, "Per-field confidence scores\nLow-confidence flagged for human review", size=7, color="#7c6354", anchor="lm")
    c.label(360, 268, "OCR artifacts corrected;\ndates standardized to ISO 8601", size=7, color="#7c6354", anchor="lm")
    c.label(360, 148, "Named entities, dates, income,\nrelationship facts extracted", size=7, color="#7c6354", anchor="lm")
    c.label(360, 68, "Validated payload enters\ndeterministic rule engine", size=7, color="#7c6354", anchor="lm")
    return c

def render_figures():
    from PIL import ImageFont  # noqa: F401
    FIG_DIR.mkdir(exist_ok=True)
    out = {}
    for name, fn in [("pipeline_tech", fig_pipeline_tech),
                     ("ai_boundary", fig_ai_boundary),
                     ("ocr_detail", fig_ocr_detail)]:
        c = fn()
        p = FIG_DIR / f"manan_{name}.png"
        c.save(str(p))
        out[name] = p
    return out

# ----------------------------------------------------------------------------
# PDF (reportlab)
# ----------------------------------------------------------------------------
def build_pdf(cover, body, figs) -> Path:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import cm
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.enums import TA_JUSTIFY, TA_CENTER, TA_LEFT
    from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer,
                                    PageBreak, Image)
    from reportlab.lib import colors
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    pdfmetrics.registerFont(TTFont("Georgia", "/System/Library/Fonts/Supplemental/Georgia.ttf"))
    pdfmetrics.registerFont(TTFont("Georgia-Bold", "/System/Library/Fonts/Supplemental/Georgia Bold.ttf"))
    pdfmetrics.registerFont(TTFont("Georgia-Italic", "/System/Library/Fonts/Supplemental/Georgia Italic.ttf"))
    pdfmetrics.registerFont(TTFont("Georgia-BoldItalic", "/System/Library/Fonts/Supplemental/Georgia Bold Italic.ttf"))
    from reportlab.pdfbase.pdfmetrics import registerFontFamily
    registerFontFamily("Georgia", normal="Georgia", bold="Georgia-Bold",
                       italic="Georgia-Italic", boldItalic="Georgia-BoldItalic")

    st_body = ParagraphStyle("body", fontName=BODY_FONT, fontSize=FONT_SIZE,
                             leading=LEADING, alignment=TA_JUSTIFY, spaceAfter=7)
    st_h2 = ParagraphStyle("h2", fontName=BODY_BOLD, fontSize=14, leading=18,
                           spaceBefore=10, spaceAfter=6, textColor=colors.black)
    st_h3 = ParagraphStyle("h3", fontName=BODY_BOLD, fontSize=12.5, leading=16,
                           spaceBefore=8, spaceAfter=5, textColor=colors.black)
    st_bullet = ParagraphStyle("bullet", parent=st_body, leftIndent=18,
                               bulletIndent=6, alignment=TA_LEFT, spaceAfter=5)
    st_cap = ParagraphStyle("caption", fontName=BODY_ITAL, fontSize=10,
                            leading=13, alignment=TA_CENTER, spaceBefore=4, spaceAfter=10,
                            textColor=colors.HexColor("#5c4a3a"))
    st_cover = ParagraphStyle("cover", fontName=BODY_FONT, fontSize=14,
                              leading=18, alignment=TA_CENTER)
    st_cover_b = ParagraphStyle("cover_b", fontName=BODY_BOLD, fontSize=13,
                                leading=17, alignment=TA_CENTER)

    def footer(canv, doc):
        if doc.page > 1:
            canv.saveState()
            canv.setFont(BODY_FONT, 9)
            canv.drawCentredString(A4[0] / 2, 1.1 * cm, f"Page {doc.page}")
            canv.restoreState()

    pdf = OUT_DIR / "manan_report.pdf"
    doc = SimpleDocTemplate(str(pdf), pagesize=A4,
                            leftMargin=2.54 * cm, rightMargin=2.54 * cm,
                            topMargin=2.54 * cm, bottomMargin=2.0 * cm,
                            title="Technical and AI Methodology Grounding of an AI-Assisted THOA Compliance System",
                            author="Manan Rastogi")
    story = []

    # ---- Cover ----
    story.append(Spacer(1, 1.2 * cm))
    if LOGO_PATH.exists():
        img = Image(str(LOGO_PATH), width=5.5 * cm, height=5.5 * cm)
        img.hAlign = "CENTER"
        story.append(img)
    story.append(Spacer(1, 1.2 * cm))
    story.append(Paragraph("<b>VELLORE INSTITUTE OF TECHNOLOGY</b>",
                           ParagraphStyle("v", parent=st_cover, fontSize=22, leading=27)))
    story.append(Spacer(1, 0.4 * cm))
    for blk in cover:
        line = fill_cover_text(blk.text)
        if "@LOGO@" in line or "VELLORE INSTITUTE OF TECHNOLOGY" in line or not line.strip():
            continue
        if "Project Title:" in line:
            story.append(Spacer(1, 4.0 * cm))
            story.append(Paragraph(html_escape(line),
                                   ParagraphStyle("c2", parent=st_cover, fontSize=15, leading=20, spaceBefore=4)))
            story.append(Spacer(1, 4.0 * cm))
        elif "Submitted by:" in line or "Submitted to:" in line:
            story.append(Paragraph(html_escape(line), st_cover_b))
        else:
            story.append(Paragraph(html_escape(line), st_cover))
    story.append(PageBreak())

    # ---- Body ----
    i = 0
    while i < len(body):
        blk = body[i]
        if blk.kind == "brk":
            story.append(PageBreak())
        elif blk.kind == "fig":
            png = figs.get(blk.text)
            if png:
                img = Image(str(png), width=9.0 * cm, height=9.0 * cm * 0.68)
                img.hAlign = "CENTER"
                story.append(img)
            i += 1
            if i < len(body) and body[i].kind == "body" and body[i].text.startswith("Figure"):
                story.append(Paragraph(html_escape(body[i].text), st_cap))
        elif blk.kind == "h2":
            story.append(Paragraph(html_escape(blk.text), st_h2))
        elif blk.kind == "h3":
            story.append(Paragraph(html_escape(blk.text), st_h3))
        elif blk.kind == "bullet":
            story.append(Paragraph(html_escape(blk.text), st_bullet, bulletText="\u2022"))
        else:
            story.append(Paragraph(html_escape(blk.text), st_body))
        i += 1

    doc.build(story, onFirstPage=lambda c, d: None, onLaterPages=footer)
    return pdf

# ----------------------------------------------------------------------------
# DOCX (python-docx)
# ----------------------------------------------------------------------------
def build_docx(cover, body, figs) -> Path:
    from docx import Document
    from docx.shared import Pt, Cm, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    doc = Document()

    normal = doc.styles["Normal"]
    normal.font.name = "Georgia"
    normal.font.size = Pt(FONT_SIZE)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Georgia")
    pf = normal.paragraph_format
    pf.line_spacing = 1.18
    pf.space_after = Pt(6)

    for name, size in (("Heading 1", 18), ("Heading 2", 14), ("Heading 3", 12.5)):
        st = doc.styles[name]
        st.font.name = "Georgia"
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.paragraph_format.line_spacing = 1.18
        st.paragraph_format.space_before = Pt(12 if name != "Heading 1" else 0)
        st.paragraph_format.space_after = Pt(6)

    sec = doc.sections[0]
    sec.top_margin = Cm(2.54)
    sec.bottom_margin = Cm(2.0)
    sec.left_margin = Cm(2.54)
    sec.right_margin = Cm(2.54)
    sec.different_first_page_header_footer = True
    ftr = sec.footer
    fp = ftr.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = fp.add_run()
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    run._r.append(fld)

    def add_rich(p, text):
        for part in re.split(r"(\*\*.*?\*\*)", text):
            if not part:
                continue
            if part.startswith("**") and part.endswith("**"):
                r = p.add_run(part[2:-2])
                r.bold = True
            else:
                p.add_run(part)

    def cover_para(text="", size=13, bold=False, space_after=6):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.15
        if text:
            r = p.add_run(text)
            r.font.size = Pt(size)
            r.bold = bold
        return p

    # ---- Cover ----
    if LOGO_PATH.exists():
        p = cover_para(space_after=18)
        p.add_run().add_picture(str(LOGO_PATH), width=Cm(5.5))
    else:
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.cell(0, 0)
        cp = cell.paragraphs[0]
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = cp.add_run("[ INSERT VIT LOGO HERE ]")
        r.font.size = Pt(9)
        r.italic = True
        r.font.color.rgb = RGBColor(128, 128, 128)

    cover_para("VELLORE INSTITUTE OF TECHNOLOGY", size=22, bold=True, space_after=22)
    for blk in cover:
        line = fill_cover_text(blk.text)
        if "@LOGO@" in line or not line.strip():
            continue
        if "VELLORE INSTITUTE OF TECHNOLOGY" in line:
            continue
        if "Project Title:" in line:
            cover_para("", space_after=24)
            cover_para(line, size=16, bold=True, space_after=26)
        elif "Submitted by:" in line or "Submitted to:" in line:
            cover_para(line, size=13, space_after=8)
        else:
            cover_para(line, size=13, space_after=6)

    doc.add_page_break()

    # ---- Body ----
    i = 0
    while i < len(body):
        blk = body[i]
        if blk.kind == "brk":
            doc.add_page_break()
        elif blk.kind == "fig":
            png = figs.get(blk.text)
            if png:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(str(png), width=Cm(9.0))
            i += 1
            if i < len(body) and body[i].kind == "body" and body[i].text.startswith("Figure"):
                cap = doc.add_paragraph()
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                cr = cap.add_run(body[i].text.replace("*", ""))
                cr.italic = True
                cr.font.size = Pt(10)
        elif blk.kind == "h2":
            p = doc.add_heading(level=2)
            add_rich(p, blk.text)
        elif blk.kind == "h3":
            p = doc.add_heading(level=3)
            add_rich(p, blk.text)
        elif blk.kind == "bullet":
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_after = Pt(5)
            add_rich(p, blk.text)
        else:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            add_rich(p, blk.text)
        i += 1

    out = OUT_DIR / "manan_report.docx"
    doc.save(str(out))
    return out

def main():
    OUT_DIR.mkdir(exist_ok=True)
    cover, body = parse_md()
    figs = render_figures()
    pdf = build_pdf(cover, body, figs)
    docx = build_docx(cover, body, figs)

    from pypdf import PdfReader
    r = PdfReader(str(pdf))
    total = len(r.pages)
    inference_page = None
    for i, page in enumerate(r.pages, 1):
        text = page.extract_text() or ""
        if "7. Inference" in text:
            inference_page = i
    print(f"PDF: {pdf} | {total} pages | Inference on page {inference_page}")
    print(f"DOCX: {docx}")
    if total != 11 or inference_page != 11:
        print("WARNING: layout is not cover=1 ... inference=11")
    return 0

if __name__ == "__main__":
    raise SystemExit(main())
